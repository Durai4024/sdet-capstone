import pytest
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
import openpyxl

# Fixture to set up the WebDriver
@pytest.fixture(scope="module")
def setup():
    # Setup the WebDriver
    driver = webdriver.Chrome()
    driver.maximize_window()
    driver.get("https://www.saucedemo.com/")
    
    yield driver
    
    # Tear down
    driver.quit()

# Fixture to read data from Excel
@pytest.fixture(scope="module")
def login_data():
    # Load Excel file
    workbook = openpyxl.load_workbook('data/login_data.xlsx')
    
    # Read valid data
    valid_data = []
    sheet = workbook['valid_data']
    for row in range(2, sheet.max_row + 1):
        username = sheet.cell(row=row, column=1).value
        password = sheet.cell(row=row, column=2).value
        valid_data.append((username, password))
    
    return valid_data
